"""
Command-line example:

#python attia_ecg_training_v2.py
--demograph-path ~/path/ptb/ptb_scp_diag_simplified.csv
--data-path ~/path/ptb/ptb-xl-a-large-publicly-available-electrocardiography-dataset-1.0.1
--batch-size 100
--max-epochs 100
--learning-rate 1e-3
--weight-decay 0.0
--patience 8
--device cuda:2
--checkpoint-dir ~/path/attia_checkpoints
--result-path ~/path/attia_kfold_results.pt
--excel-result-path ~/path/attia_kfold_testing_results.xlsx
--positive-label 1
"""

from __future__ import annotations

import copy
import os
import random
from concurrent.futures import (
    ThreadPoolExecutor,
    ProcessPoolExecutor,
)
from multiprocessing import get_context
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
import torch
import torch.nn as nn
import torch.nn.functional as F
import wfdb
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.metrics import (
    accuracy_score,
    confusion_matrix,
    f1_score,
    roc_auc_score,
    roc_curve,
)
from sklearn.model_selection import StratifiedKFold
from torch.utils.data import DataLoader, TensorDataset

# The denoiser must be available in the same project/environment.
# Change this import to the actual module path used in your project.
try:
    from ecg_denoiser import ECGWaveletKalmanBilateralDenoiser
except ImportError as exc:
    ECGWaveletKalmanBilateralDenoiser = None
    _DENOISER_IMPORT_ERROR = exc


def set_seed(seed: int = 42) -> None:
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)
    torch.cuda.manual_seed_all(seed)


##------------------------------------
# Data manager
##------------------------------------
class datamanager:
    _DENOISER_CONFIG = {
        "sampling_rate_hz": 100.0,
        "wavelet": "sym4",
        "level": 4,
        "baseline_short_window_s": 0.20,
        "baseline_long_window_s": 0.60,
        "edge_padding_s": 1.50,
        "coefficient_protect_low_sigma": 2.5,
        "coefficient_protect_high_sigma": 5.0,
        "coefficient_protection_strength": 0.30,
        "qrs_gradient_sigma": 4.0,
        "qrs_protection_radius_s": 0.055,
        "qrs_protection_by_level": (0.40, 0.55, 0.75, 0.90),
        "bilateral_radius_s": 0.030,
        "bilateral_sigma_spatial_s": 0.015,
        "bilateral_range_noise_scale": 2.5,
        "shrinkage_threshold_scale": 0.30,
        "shrinkage_strength": 0.12,
        "center_output": True,
    }

    def __init__(
        self,
        demograph_path="~/path/ptb/ptb_scp_diag_simplifed.csv",
        data_path="~/path/ptb/ptb-xl-a-large-publicly-available-electrocardiography-dataset-1.0.1",
        kfold=5,
    ):
        self.demograph_path = demograph_path
        self.data_path = data_path
        self.kfold = kfold
        self.kfold_scheme = []

    @staticmethod
    def _create_denoiser():
        return ECGWaveletKalmanBilateralDenoiser(
            **datamanager._DENOISER_CONFIG
        )

    @staticmethod
    def _denoise_one_case_worker(dct):
        ecg = np.asarray(
            dct["leads"],
            dtype=np.float32,
        )

        denoiser = datamanager._create_denoiser()

        denoised_ecg = np.empty_like(
            ecg,
            dtype=np.float32,
        )

        for lead_idx in range(ecg.shape[1]):
            denoised_ecg[:, lead_idx] = denoiser.denoise(
                ecg[:, lead_idx]
            )

        out = dict(dct)
        out["leads"] = denoised_ecg
        return out

    def load_demo(self):
        if not os.path.exists(self.demograph_path):
            return []

        df = pd.read_csv(self.demograph_path)

        columns_to_extract = [
            "ecg_id",
            "patient_id",
            "age",
            "sex",
            "weight",
            "diagnostic_class",
            "diagnostic_class_id",
            "filename_lr",
        ]

        return df[columns_to_extract]

    ##------------------------------------
    # Stage 1: Load raw ECG
    ##------------------------------------
    def load_one_case(self, dct):
        filepath_relative = dct["filename_lr"]
        filepath = os.path.join(
            self.data_path,
            filepath_relative,
        )

        record = wfdb.rdrecord(filepath)
        ecg = np.asarray(
            record.p_signal,
            dtype=np.float32,
        )

        out = dict(dct)
        out["leads"] = ecg
        return out

    def load_dataset_serial(self):
        demo = self.load_demo()

        if isinstance(demo, list) or len(demo) == 0:
            return []

        dctlist = demo.to_dict(orient="records")
        dlen = len(dctlist)
        dctset = []

        for i, idct in enumerate(dctlist):
            indct = self.load_one_case(idct)
            print(
                f"{'-' * 20} "
                f"Loading {i + 1}/{dlen}, "
                f"{indct['filename_lr']}"
            )
            dctset.append(indct)

        return dctset

    def load_dataset_parallel(self, num_workers=16):
        demo = self.load_demo()

        if isinstance(demo, list) or len(demo) == 0:
            return []

        dctlist = demo.to_dict(orient="records")

        with ThreadPoolExecutor(
            max_workers=num_workers
        ) as executor:
            dctset = list(
                executor.map(
                    self.load_one_case,
                    dctlist,
                )
            )

        return dctset

    ##------------------------------------
    # Stage 2: Denoise ECG
    ##------------------------------------
    def denoise_one_case(self, dct):
        ecg = np.asarray(
            dct["leads"],
            dtype=np.float32,
        )

        denoiser = self._create_denoiser()
        denoised_ecg = np.empty_like(
            ecg,
            dtype=np.float32,
        )

        for lead_idx in range(ecg.shape[1]):
            denoised_ecg[:, lead_idx] = denoiser.denoise(
                ecg[:, lead_idx]
            )

        out = dict(dct)
        out["leads"] = denoised_ecg
        return out

    def denoise_dataset_serial(self, dctset):
        if len(dctset) == 0:
            return []

        dctset_denoised = []
        dlen = len(dctset)

        for i, idct in enumerate(dctset):
            odct = self.denoise_one_case(idct)
            print(
                f"{'-' * 20} "
                f"Denoising {i + 1}/{dlen}, "
                f"{odct['filename_lr']}"
            )
            dctset_denoised.append(odct)

        return dctset_denoised

    def denoise_dataset_parallel(
        self,
        dctset,
        num_workers=16,
        chunksize=1,
    ):
        if len(dctset) == 0:
            return []

        context = get_context("spawn")

        with ProcessPoolExecutor(
            max_workers=num_workers,
            mp_context=context,
        ) as executor:
            dctset_denoised = list(
                executor.map(
                    datamanager._denoise_one_case_worker,
                    dctset,
                    chunksize=chunksize,
                )
            )

        return dctset_denoised

    ##------------------------------------
    # Stage 3: Extract model data
    ##------------------------------------
    def extract_model_data(self, dctset=None):
        if not dctset:
            return []

        dctlist = []

        for idct in dctset:
            dctlist.append(
                {
                    "ecg_id": idct["ecg_id"],
                    "diagnostic_class_id": idct[
                        "diagnostic_class_id"
                    ],
                    "leads": idct["leads"],
                }
            )

        return dctlist

    ##------------------------------------
    # Stage 4: K-fold split
    ##------------------------------------
    def split_kfold(
        self,
        dctlist=None,
        p1=6,
        p2=2,
        p3=2,
        n_splits=5,
        random_state=42,
    ):
        if not dctlist:
            self.kfold_scheme = []
            return False

        if not all(
            isinstance(p, int) and p > 0
            for p in (p1, p2, p3)
        ):
            raise ValueError(
                "p1, p2, and p3 must be positive integers."
            )

        if not isinstance(n_splits, int) or n_splits < 2:
            raise ValueError(
                "n_splits must be an integer >= 2."
            )

        ratio_sum = p1 + p2 + p3
        expected_test_ratio = 1.0 / n_splits
        actual_test_ratio = p3 / ratio_sum

        if not np.isclose(
            actual_test_ratio,
            expected_test_ratio,
        ):
            raise ValueError(
                "The requested ratio is not compatible with "
                "standard K-fold cross-validation. "
                f"Requested ratio: {p1}:{p2}:{p3}, "
                f"n_splits: {n_splits}. "
                f"Testing ratio must be 1/{n_splits}."
            )

        labels = np.asarray(
            [
                dct["diagnostic_class_id"]
                for dct in dctlist
            ]
        )

        indices = np.arange(len(dctlist))

        outer_kfold = StratifiedKFold(
            n_splits=n_splits,
            shuffle=True,
            random_state=random_state,
        )

        fold_list = []

        for fold_idx, (remaining_idx, test_idx) in enumerate(
            outer_kfold.split(indices, labels)
        ):
            remaining_labels = labels[remaining_idx]

            # For 6:2:2, the remaining 80% is split into 75% training
            # and 25% validation, yielding 60% and 20% overall.
            train_val_kfold = StratifiedKFold(
                n_splits=4,
                shuffle=True,
                random_state=random_state + fold_idx + 1,
            )

            train_relative_idx, val_relative_idx = next(
                train_val_kfold.split(
                    remaining_idx,
                    remaining_labels,
                )
            )

            train_idx = remaining_idx[train_relative_idx]
            val_idx = remaining_idx[val_relative_idx]

            training_list = [dctlist[i] for i in train_idx]
            validation_list = [dctlist[i] for i in val_idx]
            testing_list = [dctlist[i] for i in test_idx]

            training = self._merge_split(training_list)
            validation = self._merge_split(validation_list)
            testing = self._merge_split(testing_list)

            fold_list.append(
                {
                    "fold": fold_idx,
                    "training": training,
                    "validation": validation,
                    "testing": testing,
                }
            )

        self.kfold_scheme = fold_list
        return True

    def _merge_split(self, dctlist):
        if len(dctlist) == 0:
            return {
                "leads": torch.empty((0, 0, 0), dtype=torch.float32),
                "diagnostic_class_id": [],
                "ecg_id": [],
            }

        leads = np.asarray(
            [dct["leads"] for dct in dctlist],
            dtype=np.float32,
        )

        return {
            "leads": torch.tensor(
                leads,
                dtype=torch.float32,
            ),
            "diagnostic_class_id": [
                dct["diagnostic_class_id"]
                for dct in dctlist
            ],
            "ecg_id": [
                dct["ecg_id"]
                for dct in dctlist
            ],
        }

    ##------------------------------------
    # Optional persistence
    ##------------------------------------
    @staticmethod
    def save_dctset(dctset, filepath):
        os.makedirs(
            os.path.dirname(filepath) or ".",
            exist_ok=True,
        )
        torch.save(dctset, filepath)

    @staticmethod
    def load_dctset(filepath):
        return torch.load(
            filepath,
            map_location="cpu",
            weights_only=False,
        )

    ##------------------------------------
    # Full pipeline
    ##------------------------------------
    def prepare_from_raw(
        self,
        load_parallel=True,
        load_workers=16,
        denoise_parallel=True,
        denoise_workers=16,
        denoise_chunksize=1,
        split_random_state=0,
    ):
        """Load raw ECG, denoise, extract model data, and create k-fold scheme."""
        if load_parallel:
            raw_dctset = self.load_dataset_parallel(
                num_workers=load_workers
            )
        else:
            raw_dctset = self.load_dataset_serial()

        if len(raw_dctset) == 0:
            self.kfold_scheme = []
            return False

        if denoise_parallel:
            denoised_dctset = self.denoise_dataset_parallel(
                dctset=raw_dctset,
                num_workers=denoise_workers,
                chunksize=denoise_chunksize,
            )
        else:
            denoised_dctset = self.denoise_dataset_serial(
                dctset=raw_dctset
            )

        dctlist = self.extract_model_data(
            dctset=denoised_dctset
        )

        return self.split_kfold(
            dctlist=dctlist,
            p1=6,
            p2=2,
            p3=2,
            n_splits=self.kfold,
            random_state=split_random_state,
        )

    def prepare_from_dctset(
        self,
        dctset,
        random_state=0,
    ):
        """Create k-fold scheme from an already prepared denoised dctset."""
        dctlist = self.extract_model_data(
            dctset=dctset
        )

        return self.split_kfold(
            dctlist=dctlist,
            p1=6,
            p2=2,
            p3=2,
            n_splits=self.kfold,
            random_state=random_state,
        )

    def run(
        self,
        load_parallel=True,
        load_workers=16,
        denoise_parallel=True,
        denoise_workers=16,
        denoise_chunksize=1,
        random_state=0,
    ):
        """Run the complete data preparation pipeline explicitly."""
        print("=" * 80)
        print("Step 1/4: Loading raw ECG data")
        print("=" * 80)

        if load_parallel:
            raw_dctset = self.load_dataset_parallel(
                num_workers=load_workers
            )
        else:
            raw_dctset = self.load_dataset_serial()

        if len(raw_dctset) == 0:
            self.kfold_scheme = []
            print("No ECG data were loaded.")
            return self.kfold_scheme

        print(f"Loaded {len(raw_dctset)} ECG records.")

        print("=" * 80)
        print("Step 2/4: Denoising ECG data")
        print("=" * 80)

        if denoise_parallel:
            denoised_dctset = self.denoise_dataset_parallel(
                dctset=raw_dctset,
                num_workers=denoise_workers,
                chunksize=denoise_chunksize,
            )
        else:
            denoised_dctset = self.denoise_dataset_serial(
                dctset=raw_dctset
            )

        if len(denoised_dctset) == 0:
            self.kfold_scheme = []
            print("No ECG data remain after denoising.")
            return self.kfold_scheme

        print(f"Denoised {len(denoised_dctset)} ECG records.")

        print("=" * 80)
        print("Step 3/4: Extracting model data")
        print("=" * 80)

        dctlist = self.extract_model_data(
            dctset=denoised_dctset
        )

        if len(dctlist) == 0:
            self.kfold_scheme = []
            print("No model data were extracted.")
            return self.kfold_scheme

        print(f"Prepared {len(dctlist)} model records.")

        print("=" * 80)
        print("Step 4/4: Creating 6:2:2 k-fold scheme")
        print("=" * 80)

        success = self.split_kfold(
            dctlist=dctlist,
            p1=6,
            p2=2,
            p3=2,
            n_splits=self.kfold,
            random_state=random_state,
        )

        if not success:
            self.kfold_scheme = []
            print("K-fold splitting failed.")
            return self.kfold_scheme

        print(f"Generated {len(self.kfold_scheme)} folds.")

        for fold in self.kfold_scheme:
            print(
                f"Fold {fold['fold']}: "
                f"training={fold['training']['leads'].shape}, "
                f"validation={fold['validation']['leads'].shape}, "
                f"testing={fold['testing']['leads'].shape}"
            )

        return self.kfold_scheme

##------------------------------------
# Attia residual CNN
##------------------------------------

class ResidualBlock(nn.Module):
    def __init__(self, in_channels: int, out_channels: int, kernel_size: int):
        super().__init__()

        if kernel_size % 2 == 0:
            raise ValueError("kernel_size must be odd")

        padding = kernel_size // 2

        self.bn1 = nn.BatchNorm2d(in_channels)
        self.relu1 = nn.ReLU(inplace=True)
        self.conv1 = nn.Conv2d(
            in_channels,
            out_channels,
            kernel_size=(kernel_size, 1),
            padding=(padding, 0),
            bias=False,
        )

        self.bn2 = nn.BatchNorm2d(out_channels)
        self.relu2 = nn.ReLU(inplace=True)
        self.conv2 = nn.Conv2d(
            out_channels,
            out_channels,
            kernel_size=(kernel_size, 1),
            padding=(padding, 0),
            bias=False,
        )

        if in_channels != out_channels:
            self.shortcut = nn.Conv2d(
                in_channels,
                out_channels,
                kernel_size=1,
                bias=False,
            )
        else:
            self.shortcut = nn.Identity()

        self.pool = nn.MaxPool2d(kernel_size=(2, 1), stride=(2, 1))

    def forward(self, x):
        identity = self.shortcut(x)

        out = self.bn1(x)
        out = self.relu1(out)
        out = self.conv1(out)

        out = self.bn2(out)
        out = self.relu2(out)
        out = self.conv2(out)

        out = out + identity
        return self.pool(out)


class AttiaECGNet(nn.Module):
    PAPER_LEAD_INDICES = (0, 1, 6, 7, 8, 9, 10, 11)

    def __init__(self, dropout=0.2, num_classes=2):
        super().__init__()
        if num_classes != 2:
            raise ValueError("This implementation is binary classification only.")

        self.input_samples = 1000
        self.input_leads = 12
        self.padded_samples = 1024

        block_config = [
            (7, 16), (7, 16), (7, 16),
            (5, 32), (5, 32), (5, 32),
            (3, 64), (3, 64), (3, 64),
        ]

        layers = []
        in_channels = 1
        for kernel_size, out_channels in block_config:
            layers.append(
                ResidualBlock(
                    in_channels=in_channels,
                    out_channels=out_channels,
                    kernel_size=kernel_size,
                )
            )
            in_channels = out_channels

        self.residual_layers = nn.Sequential(*layers)
        self.fusion = nn.Conv2d(
            64,
            128,
            kernel_size=(1, 8),
            bias=False,
        )
        self.fusion_bn = nn.BatchNorm2d(128)
        self.fusion_relu = nn.ReLU(inplace=True)
        self.dropout = nn.Dropout(dropout)

        self.output_time = self.padded_samples // (2 ** len(block_config))
        self.classifier = nn.Linear(128 * self.output_time, num_classes)
        self._initialize_weights()

    def _initialize_weights(self):
        for module in self.modules():
            if isinstance(module, nn.Conv2d):
                nn.init.kaiming_normal_(
                    module.weight,
                    mode="fan_out",
                    nonlinearity="relu",
                )
            elif isinstance(module, nn.Linear):
                nn.init.kaiming_normal_(
                    module.weight,
                    mode="fan_in",
                    nonlinearity="linear",
                )
                nn.init.zeros_(module.bias)
            elif isinstance(module, nn.BatchNorm2d):
                nn.init.ones_(module.weight)
                nn.init.zeros_(module.bias)

    def forward(self, x):
        if x.ndim != 3 or x.shape[1:] != (1000, 12):
            raise ValueError(
                f"Expected input [B, 1000, 12], got {tuple(x.shape)}"
            )

        x = x.float()
        indices = torch.as_tensor(
            self.PAPER_LEAD_INDICES,
            device=x.device,
            dtype=torch.long,
        )
        x = x.index_select(dim=-1, index=indices)

        x = F.pad(
            x,
            pad=(0, 0, 0, self.padded_samples - self.input_samples),
            mode="constant",
            value=0.0,
        )
        x = x.unsqueeze(1)
        x = self.residual_layers(x)
        x = self.fusion(x)
        x = self.fusion_bn(x)
        x = self.fusion_relu(x)
        x = self.dropout(x)
        x = torch.flatten(x, start_dim=1)
        return self.classifier(x)


##------------------------------------
# Label encoding and metrics
##------------------------------------

def encode_binary_labels(labels, positive_label=None):
    labels = list(labels)
    unique = list(dict.fromkeys(labels))

    if len(unique) != 2:
        raise ValueError(
            "The current network is binary. "
            f"Expected two diagnostic_class_id values, got {unique}."
        )

    if positive_label is None:
        ordered = sorted(unique, key=lambda x: str(x))
        mapping = {ordered[0]: 0, ordered[1]: 1}
    else:
        if positive_label not in unique:
            raise ValueError(
                f"positive_label={positive_label!r} is not in {unique}."
            )
        negative_label = next(x for x in unique if x != positive_label)
        mapping = {negative_label: 0, positive_label: 1}

    encoded = np.asarray([mapping[x] for x in labels], dtype=np.int64)
    return encoded, mapping


def safe_auc(y_true, y_prob):
    if len(np.unique(y_true)) < 2:
        return float("nan")
    return float(roc_auc_score(y_true, y_prob))


def binary_metrics(y_true, y_prob, threshold):
    y_pred = (y_prob >= threshold).astype(np.int64)
    tn, fp, fn, tp = confusion_matrix(
        y_true,
        y_pred,
        labels=[0, 1],
    ).ravel()

    sensitivity = tp / (tp + fn) if tp + fn > 0 else 0.0
    specificity = tn / (tn + fp) if tn + fp > 0 else 0.0

    return {
        "auc": safe_auc(y_true, y_prob),
        "accuracy": float(accuracy_score(y_true, y_pred)),
        "sensitivity": float(sensitivity),
        "specificity": float(specificity),
        "f1": float(f1_score(y_true, y_pred, zero_division=0)),
        "threshold": float(threshold),
    }


def select_validation_threshold(y_true, y_prob):
    fpr, tpr, thresholds = roc_curve(y_true, y_prob)
    scores = 0.5 * (tpr + (1.0 - fpr))
    finite = np.isfinite(thresholds)
    if not finite.any():
        return 0.5

    scores = scores.copy()
    scores[~finite] = -np.inf
    index = int(np.argmax(scores))
    threshold = float(thresholds[index])

    if not np.isfinite(threshold):
        threshold = 0.5
    return threshold


##------------------------------------
# Training configuration and trainer
##------------------------------------

@dataclass
class TrainConfig:
    batch_size: int = 32
    max_epochs: int = 50
    learning_rate: float = 1e-3
    weight_decay: float = 0.0
    patience: int = 8
    num_workers: int = 0


class ECGTrainer:
    def __init__(self, model, config, device=None):
        self.model = model
        self.config = config
        self.device = torch.device(
            device if device is not None else (
                "cuda" if torch.cuda.is_available() else "cpu"
            )
        )
        self.model.to(self.device)
        self.criterion = nn.CrossEntropyLoss()
        self.optimizer = torch.optim.Adam(
            self.model.parameters(),
            lr=config.learning_rate,
            weight_decay=config.weight_decay,
        )
        self.best_epoch = -1
        self.best_threshold = 0.5
        self.history = []

    @staticmethod
    def make_loader(split, labels, batch_size, shuffle, num_workers, device):
        leads = split["leads"]
        if not torch.is_tensor(leads):
            leads = torch.as_tensor(leads)
        leads = leads.float().contiguous()
        labels = torch.as_tensor(labels, dtype=torch.long)

        dataset = TensorDataset(leads, labels)
        return DataLoader(
            dataset,
            batch_size=batch_size,
            shuffle=shuffle,
            num_workers=num_workers,
            pin_memory=device.type == "cuda",
        )

    def run_epoch(self, loader, training):
        self.model.train(training)
        total_loss = 0.0
        true_list = []
        prob_list = []

        for x, y in loader:
            x = x.to(self.device, non_blocking=True)
            y = y.to(self.device, non_blocking=True)

            with torch.set_grad_enabled(training):
                logits = self.model(x)
                loss = self.criterion(logits, y)

                if training:
                    self.optimizer.zero_grad(set_to_none=True)
                    loss.backward()
                    self.optimizer.step()

            total_loss += float(loss.item()) * x.shape[0]
            true_list.append(y.detach().cpu().numpy())
            prob_list.append(torch.softmax(logits.detach(), dim=1)[:, 1].cpu().numpy())

        return (
            total_loss / len(loader.dataset),
            np.concatenate(true_list),
            np.concatenate(prob_list),
        )

    def fit(
        self,
        training_split,
        training_labels,
        validation_split,
        validation_labels,
        checkpoint_path=None,
    ):
        train_loader = self.make_loader(
            training_split,
            training_labels,
            self.config.batch_size,
            True,
            self.config.num_workers,
            self.device,
        )
        val_loader = self.make_loader(
            validation_split,
            validation_labels,
            self.config.batch_size,
            False,
            self.config.num_workers,
            self.device,
        )

        best_val_auc = -np.inf
        best_state = None
        bad_epochs = 0

        for epoch in range(1, self.config.max_epochs + 1):
            train_loss, train_true, train_prob = self.run_epoch(train_loader, True)
            val_loss, val_true, val_prob = self.run_epoch(val_loader, False)

            train_auc = safe_auc(train_true, train_prob)
            val_auc = safe_auc(val_true, val_prob)

            self.history.append({
                "epoch": epoch,
                "train_loss": train_loss,
                "val_loss": val_loss,
                "train_auc": train_auc,
                "val_auc": val_auc,
            })

            print(
                f"Epoch {epoch:03d} | "
                f"train_loss={train_loss:.5f} | "
                f"val_loss={val_loss:.5f} | "
                f"train_auc={train_auc:.4f} | "
                f"val_auc={val_auc:.4f}"
            )

            if np.isfinite(val_auc) and val_auc > best_val_auc:
                best_val_auc = val_auc
                best_state = copy.deepcopy(self.model.state_dict())
                self.best_epoch = epoch
                bad_epochs = 0
            else:
                bad_epochs += 1

            if bad_epochs >= self.config.patience:
                print(f"Early stopping at epoch {epoch}.")
                break

        if best_state is None:
            raise RuntimeError("No valid validation checkpoint was obtained.")

        self.model.load_state_dict(best_state)

        _, val_true, val_prob = self.run_epoch(val_loader, False)
        self.best_threshold = select_validation_threshold(val_true, val_prob)

        if checkpoint_path is not None:
            checkpoint_path = Path(checkpoint_path)
            checkpoint_path.parent.mkdir(parents=True, exist_ok=True)
            torch.save(
                {
                    "model_state_dict": self.model.state_dict(),
                    "optimizer_state_dict": self.optimizer.state_dict(),
                    "best_epoch": self.best_epoch,
                    "best_val_auc": best_val_auc,
                    "best_threshold": self.best_threshold,
                    "history": self.history,
                },
                checkpoint_path,
            )

    @torch.no_grad()
    def predict(self, split):
        leads = split["leads"]
        if not torch.is_tensor(leads):
            leads = torch.as_tensor(leads)

        dummy = np.zeros(leads.shape[0], dtype=np.int64)
        loader = self.make_loader(
            split,
            dummy,
            self.config.batch_size,
            False,
            self.config.num_workers,
            self.device,
        )

        self.model.eval()
        probabilities = []
        for x, _ in loader:
            x = x.to(self.device, non_blocking=True)
            logits = self.model(x)
            probabilities.append(torch.softmax(logits, dim=1)[:, 1].cpu().numpy())

        y_prob = np.concatenate(probabilities)
        y_pred = (y_prob >= self.best_threshold).astype(np.int64)
        return y_prob, y_pred

    @torch.no_grad()
    def evaluate(self, split, labels, split_name):
        y_prob, y_pred = self.predict(split)
        result = binary_metrics(labels, y_prob, self.best_threshold)
        result["split"] = split_name
        result["predicted_positive"] = int(np.sum(y_pred == 1))
        result["n_samples"] = int(len(labels))
        return result, y_prob, y_pred

##------------------------------------
# Statistical analysis across folds
##------------------------------------

def summarize_fold_results(fold_results, confidence_level=0.95):
    metrics = ["auc", "accuracy", "sensitivity", "specificity", "f1"]
    alpha = 1.0 - confidence_level
    summary = {}

    for metric in metrics:
        values = np.asarray(
            [item["testing"][metric] for item in fold_results],
            dtype=float,
        )
        mean_value = float(np.nanmean(values))
        std_value = float(np.nanstd(values, ddof=1))
        n = int(np.sum(np.isfinite(values)))

        if n > 1:
            standard_error = std_value / np.sqrt(n)
            z = 1.959963984540054
            ci_low = mean_value - z * standard_error
            ci_high = mean_value + z * standard_error
        else:
            ci_low = float("nan")
            ci_high = float("nan")

        summary[metric] = {
            "mean": mean_value,
            "std": std_value,
            "ci95_low": float(ci_low),
            "ci95_high": float(ci_high),
            "values": values.tolist(),
        }

    print("\n" + "=" * 80)
    print("5-fold testing statistical summary")
    print("=" * 80)
    print(f"Confidence level: {confidence_level:.2f}")

    for metric in metrics:
        item = summary[metric]
        print(
            f"{metric:12s}: "
            f"mean={item['mean']:.4f}, "
            f"std={item['std']:.4f}, "
            f"95% CI=({item['ci95_low']:.4f}, {item['ci95_high']:.4f})"
        )

    return summary

##------------------------------------
# Export testing results to Excel
##------------------------------------

def export_testing_results_to_excel(
    fold_results,
    excel_path="~/path/attia_kfold_testing_results.xlsx",
):
    """Export all testing ECG results to five Excel worksheets, one per fold."""
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    headers = [
        "ecg_id",
        "true_label",
        "probability",
        "threshold",
        "prediction",
    ]

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_font = Font(bold=True)

    for fold_result in fold_results:
        fold_id = int(fold_result["fold"])
        worksheet = workbook.create_sheet(title=f"Fold_{fold_id}")

        testing_results = fold_result["testing_results"]
        last_row = len(testing_results) + 1
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:E{last_row}"

        for column_index, header in enumerate(headers, start=1):
            cell = worksheet.cell(row=1, column=column_index, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row_index, result in enumerate(testing_results, start=2):
            worksheet.cell(row=row_index, column=1, value=result["ecg_id"])
            worksheet.cell(row=row_index, column=2, value=result["true_label"])
            worksheet.cell(row=row_index, column=3, value=result["probability"])
            worksheet.cell(row=row_index, column=4, value=result["threshold"])
            worksheet.cell(row=row_index, column=5, value=result["prediction"])

        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
            row[0].alignment = Alignment(horizontal="center")
            row[1].alignment = Alignment(horizontal="center")
            row[2].number_format = "0.000000"
            row[3].number_format = "0.000000"
            row[4].alignment = Alignment(horizontal="center")

        widths = {1: 16, 2: 14, 3: 16, 4: 16, 5: 14}
        for column_index, width in widths.items():
            worksheet.column_dimensions[get_column_letter(column_index)].width = width

    output_path = Path(excel_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    print(f"\nSaved all five-fold testing results to: {output_path}")
    return str(output_path)

##------------------------------------
# Full five-fold experiment
##------------------------------------

def train_and_test_datamanager(
    data_manager,
    positive_label=None,
    train_config=None,
    device=None,
    checkpoint_dir="attia_checkpoints",
    result_path="attia_kfold_results.pt",
):
    if not hasattr(data_manager, "kfold_scheme"):
        raise AttributeError("The supplied object has no kfold_scheme attribute.")

    kfold_scheme = data_manager.kfold_scheme
    if not isinstance(kfold_scheme, list) or len(kfold_scheme) == 0:
        raise ValueError("data_manager.kfold_scheme is empty or invalid.")

    config = train_config or TrainConfig()
    results = []

    for fold in kfold_scheme:
        fold_id = int(fold["fold"])
        training = fold["training"]
        validation = fold["validation"]
        testing = fold["testing"]

        print("\n" + "#" * 80)
        print(f"FOLD {fold_id}")
        print("#" * 80)
        print(
            f"Training:   {training['leads'].shape[0]} ECGs"
        )
        print(
            f"Validation: {validation['leads'].shape[0]} ECGs"
        )
        print(
            f"Testing:    {testing['leads'].shape[0]} ECGs"
        )

        all_labels = (
            list(training["diagnostic_class_id"])
            + list(validation["diagnostic_class_id"])
            + list(testing["diagnostic_class_id"])
        )
        _, label_mapping = encode_binary_labels(
            all_labels,
            positive_label=positive_label,
        )

        training_labels = np.asarray(
            [label_mapping[x] for x in training["diagnostic_class_id"]],
            dtype=np.int64,
        )
        validation_labels = np.asarray(
            [label_mapping[x] for x in validation["diagnostic_class_id"]],
            dtype=np.int64,
        )
        testing_labels = np.asarray(
            [label_mapping[x] for x in testing["diagnostic_class_id"]],
            dtype=np.int64,
        )

        print(f"Label mapping: {label_mapping}")
        print(
            "Training class counts: "
            f"{dict(zip(*np.unique(training_labels, return_counts=True)))}"
        )
        print(
            "Validation class counts: "
            f"{dict(zip(*np.unique(validation_labels, return_counts=True)))}"
        )
        print(
            "Testing class counts: "
            f"{dict(zip(*np.unique(testing_labels, return_counts=True)))}"
        )

        set_seed(42 + fold_id)
        model = AttiaECGNet(dropout=0.2, num_classes=2)
        trainer = ECGTrainer(model, config, device=device)

        checkpoint_path = Path(checkpoint_dir) / f"fold_{fold_id}_best.pt"

        print("\nTraining network")
        trainer.fit(
            training_split=training,
            training_labels=training_labels,
            validation_split=validation,
            validation_labels=validation_labels,
            checkpoint_path=str(checkpoint_path),
        )

        print(
            f"Best epoch: {trainer.best_epoch}; "
            f"validation threshold: {trainer.best_threshold:.6f}"
        )

        print("\nValidation evaluation")
        validation_metrics, validation_prob, validation_pred = trainer.evaluate(
            validation,
            validation_labels,
            "validation",
        )
        print(validation_metrics)

        print("\nTesting network")
        testing_metrics, testing_prob, testing_pred = trainer.evaluate(
            testing,
            testing_labels,
            "testing",
        )
        print(testing_metrics)

        ##------------------------------------
        # Save one row per testing ECG.
        # The threshold is selected only from the validation set and is
        # fixed for every ECG in the current testing fold.
        ##------------------------------------
        testing_threshold = float(trainer.best_threshold)
        testing_results = []

        for ecg_id, true_label, probability, prediction in zip(
            testing["ecg_id"],
            testing_labels,
            testing_prob,
            testing_pred,
        ):
            testing_results.append(
                {
                    "ecg_id": ecg_id,
                    "true_label": int(true_label),
                    "probability": float(probability),
                    "threshold": testing_threshold,
                    "prediction": int(prediction),
                }
            )

        print("\nTesting prediction examples")
        print(
            "ecg_id\ttrue_label\tprobability\tthreshold\tprediction"
        )
        print("-" * 72)
        for row in testing_results[:10]:
            print(
                f"{row['ecg_id']}\t"
                f"{row['true_label']}\t"
                f"{row['probability']:.6f}\t"
                f"{row['threshold']:.6f}\t"
                f"{row['prediction']}"
            )

        results.append(
            {
                "fold": fold_id,
                "label_mapping": label_mapping,
                "best_epoch": trainer.best_epoch,
                "threshold": testing_threshold,
                "training_ecg_id": list(training["ecg_id"]),
                "validation_ecg_id": list(validation["ecg_id"]),
                "testing_ecg_id": list(testing["ecg_id"]),
                "validation": validation_metrics,
                "testing": testing_metrics,
                "validation_probability": validation_prob,
                "validation_prediction": validation_pred,
                "testing_probability": testing_prob,
                "testing_prediction": testing_pred,
                "testing_results": testing_results,
                "history": trainer.history,
            }
        )

    statistical_summary = summarize_fold_results(results)

    output = {
        "fold_results": results,
        "statistical_summary": statistical_summary,
    }

    torch.save(output, result_path)
    print(f"\nSaved all experiment results to: {result_path}")

    return output

##------------------------------------
# Main program: data loading -> split -> training -> validation -> testing -> statistics
##------------------------------------

if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(
        description=(
            "Attia ECG 5-fold training, validation, testing, "
            "checkpointing, and result export."
        )
    )

    parser.add_argument("--demograph-path", type=str, default="~/path/ptb/ptb_scp_diag_simplified.csv")
    parser.add_argument("--data-path", type=str, default="~/path/ptb/ptb-xl-a-large-publicly-available-electrocardiography-dataset-1.0.1")
    parser.add_argument("--batch-size", type=int, default=100)
    parser.add_argument("--max-epochs", type=int, default=100)
    parser.add_argument("--learning-rate", type=float, default=1e-3)
    parser.add_argument("--weight-decay", type=float, default=0.0)
    parser.add_argument("--patience", type=int, default=8)
    parser.add_argument("--device", type=str, default="cuda:2")
    parser.add_argument("--checkpoint-dir", type=str, default="~/path/attia_checkpoints")
    parser.add_argument("--result-path", type=str, default="~/path/attia_kfold_results.pt")
    parser.add_argument("--excel-result-path", type=str, default="~/path/attia_kfold_testing_results.xlsx")
    parser.add_argument("--positive-label", type=int, default=1)
    parser.add_argument("--kfold", type=int, default=5)
    parser.add_argument("--load-workers", type=int, default=64)
    parser.add_argument("--denoise-workers", type=int, default=64)
    parser.add_argument("--denoise-chunksize", type=int, default=1)
    parser.add_argument("--split-random-state", type=int, default=0)

    args = parser.parse_args()

    set_seed(42)

    if args.batch_size <= 0 or args.max_epochs <= 0 or args.patience <= 0:
        parser.error("batch-size, max-epochs, and patience must be > 0")
    if args.learning_rate <= 0:
        parser.error("learning-rate must be > 0")
    if args.weight_decay < 0:
        parser.error("weight-decay must be >= 0")
    if args.kfold < 2:
        parser.error("kfold must be >= 2")
    if args.load_workers <= 0 or args.denoise_workers <= 0 or args.denoise_chunksize <= 0:
        parser.error("worker counts and chunksize must be > 0")

    if args.device.startswith("cuda"):
        if not torch.cuda.is_available():
            raise RuntimeError("CUDA device was requested, but CUDA is not available.")
        DEVICE = torch.device(args.device)
        if DEVICE.index is not None and DEVICE.index >= torch.cuda.device_count():
            raise RuntimeError(
                f"Requested {DEVICE}, but only {torch.cuda.device_count()} CUDA device(s) are visible."
            )
        if DEVICE.index is not None:
            print(f"Using device: {DEVICE}")
            print(f"GPU: {torch.cuda.get_device_name(DEVICE.index)}")
        else:
            print(f"Using device: {DEVICE}")
    else:
        DEVICE = torch.device(args.device)
        print(f"Using device: {DEVICE}")

    DEMOGRAPH_PATH = str(Path(args.demograph_path).expanduser().resolve())
    DATA_PATH = str(Path(args.data_path).expanduser().resolve())

    if not Path(DEMOGRAPH_PATH).is_file():
        raise FileNotFoundError(
            f"Demographic CSV was not found: {DEMOGRAPH_PATH}"
        )

    if not Path(DATA_PATH).is_dir():
        raise FileNotFoundError(
            f"PTB-XL data directory was not found: {DATA_PATH}"
        )

    BATCH_SIZE = args.batch_size
    MAX_EPOCHS = args.max_epochs
    LEARNING_RATE = args.learning_rate
    WEIGHT_DECAY = args.weight_decay
    PATIENCE = args.patience
    CHECKPOINT_DIR = str(Path(args.checkpoint_dir).expanduser().resolve())
    RESULT_PATH = str(Path(args.result_path).expanduser().resolve())
    EXCEL_RESULT_PATH = str(Path(args.excel_result_path).expanduser().resolve())
    POSITIVE_LABEL = args.positive_label

    print("\n" + "=" * 80)
    print("V2 EXPERIMENT CONFIGURATION")
    print("=" * 80)
    print(f"DEMOGRAPH_PATH    = {DEMOGRAPH_PATH}")
    print(f"DATA_PATH         = {DATA_PATH}")
    print(f"BATCH_SIZE        = {BATCH_SIZE}")
    print(f"MAX_EPOCHS        = {MAX_EPOCHS}")
    print(f"LEARNING_RATE     = {LEARNING_RATE}")
    print(f"WEIGHT_DECAY      = {WEIGHT_DECAY}")
    print(f"PATIENCE          = {PATIENCE}")
    print(f"DEVICE            = {DEVICE}")
    print(f"CHECKPOINT_DIR    = {CHECKPOINT_DIR}")
    print(f"RESULT_PATH       = {RESULT_PATH}")
    print(f"EXCEL_RESULT_PATH = {EXCEL_RESULT_PATH}")
    print(f"POSITIVE_LABEL    = {POSITIVE_LABEL}")
    print(f"K-FOLD            = {args.kfold}")
    print(f"LOAD_WORKERS      = {args.load_workers}")
    print(f"DENOISE_WORKERS   = {args.denoise_workers}")
    print(f"DENOISE_CHUNKSIZE = {args.denoise_chunksize}")
    print(f"SPLIT_RANDOM_SEED = {args.split_random_state}")

    print("\n" + "=" * 80)
    print("1. DATA LOADING")
    print("=" * 80)

    dm = datamanager(
        demograph_path=DEMOGRAPH_PATH,
        data_path=DATA_PATH,
        kfold=args.kfold,
    )

    dm.run(
        load_parallel=True,
        load_workers=args.load_workers,
        denoise_parallel=True,
        denoise_workers=args.denoise_workers,
        denoise_chunksize=args.denoise_chunksize,
        random_state=args.split_random_state,
    )

    print(f"Number of folds: {len(dm.kfold_scheme)}")

    kfold_data_path = Path(CHECKPOINT_DIR).parent / "attia_kfold_scheme.pt"
    kfold_data_path.parent.mkdir(parents=True, exist_ok=True)
    torch.save(dm.kfold_scheme, kfold_data_path)
    print(f"Saved k-fold scheme to: {kfold_data_path}")

    print("\n" + "=" * 80)
    print("2. K-FOLD DATA STRUCTURE")
    print("=" * 80)
    for fold in dm.kfold_scheme:
        print(
            f"Fold {fold['fold']}: "
            f"training={tuple(fold['training']['leads'].shape)}, "
            f"validation={tuple(fold['validation']['leads'].shape)}, "
            f"testing={tuple(fold['testing']['leads'].shape)}"
        )

    train_config = TrainConfig(
        batch_size=BATCH_SIZE,
        max_epochs=MAX_EPOCHS,
        learning_rate=LEARNING_RATE,
        weight_decay=WEIGHT_DECAY,
        patience=PATIENCE,
        num_workers=0,
    )

    print("\n" + "=" * 80)
    print("3. TRAINING / VALIDATION / TESTING")
    print("=" * 80)

    results = train_and_test_datamanager(
        data_manager=dm,
        positive_label=POSITIVE_LABEL,
        train_config=train_config,
        device=DEVICE,
        checkpoint_dir=CHECKPOINT_DIR,
        result_path=RESULT_PATH,
    )

    print("\n" + "=" * 80)
    print("4. FINAL STATISTICAL ANALYSIS")
    print("=" * 80)
    for metric, values in results["statistical_summary"].items():
        print(
            f"{metric}: mean={values['mean']:.4f}, "
            f"std={values['std']:.4f}, "
            f"95% CI=({values['ci95_low']:.4f}, {values['ci95_high']:.4f})"
        )

    print("\n" + "=" * 80)
    print("5. EXPORT TESTING RESULTS TO EXCEL")
    print("=" * 80)

    export_testing_results_to_excel(
        fold_results=results["fold_results"],
        excel_path=EXCEL_RESULT_PATH,
    )

    print("\nExperiment completed.")
